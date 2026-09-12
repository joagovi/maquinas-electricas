#!/usr/bin/env python3
"""Genera la hoja de control de participacion en clase.

Una fila por alumno y una columna por cada sesion de clase del profesor
responsable hasta el parcial. Suma sola y marca en color quien no ha
participado nunca.

La salida va a privado/, que esta fuera de git: contiene nombres de alumnos.

Uso:
    .venv/bin/python scripts/participacion.py [salida.xlsx]
"""

from __future__ import annotations

import datetime
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
MESES = ["ene", "feb", "mar", "abr", "may", "jun",
         "jul", "ago", "set", "oct", "nov", "dic"]

AZUL = PatternFill("solid", fgColor="1F4E79")
GRIS = PatternFill("solid", fgColor="F2F2F2")
VERDE = PatternFill("solid", fgColor="C6EFCE")
AMBAR = PatternFill("solid", fgColor="FFF2CC")
ROJO = PatternFill("solid", fgColor="FFC7CE")
BORDE = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def fecha_de(s) -> datetime.date:
    f = s["fecha"]
    return f if isinstance(f, datetime.date) else datetime.date.fromisoformat(str(f))


def main() -> int:
    curso = yaml.safe_load((RAIZ / "silabo" / "curso.yml").read_text(encoding="utf-8"))
    alumnos = yaml.safe_load((RAIZ / "silabo" / "interno.yml").read_text(encoding="utf-8"))["alumnos"]
    resp = (curso.get("exposiciones") or {}).get("responsable", "JGV")

    # Solo las clases del profesor responsable y solo antes del parcial.
    examenes = [s["n"] for s in curso["sesiones"] if s["tipo"] == "examen"]
    tope = min(examenes) if examenes else 99
    ses = [s for s in curso["sesiones"]
           if s["tipo"] == "clase" and s.get("responsable") == resp and s["n"] < tope]

    caps = {c.get("semana"): c["titulo"] for c in curso.get("capitulos", []) if c.get("semana")}
    hoy = datetime.date.today()

    wb = Workbook()
    ws = wb.active
    ws.title = "Participación"

    ws["A1"] = f'Participación en clase · {curso["codigo"]} {curso["ciclo"]}'
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3 + len(ses))
    ws["A2"] = ("Anota un punto (o la fracción que uses) en la clase correspondiente. "
                "El total se calcula solo.")
    ws["A2"].font = Font(italic=True, size=9, color="404040")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=3 + len(ses))

    # fila 4: numero de sesion / fila 5: fecha
    ws.cell(row=5, column=1, value="Código").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=5, column=2, value="Alumno").font = Font(bold=True, color="FFFFFF")
    for c in (1, 2):
        ws.cell(row=5, column=c).fill = AZUL
        ws.cell(row=4, column=c).fill = AZUL

    for i, s in enumerate(ses):
        col = 3 + i
        f = fecha_de(s)
        tit = ws.cell(row=4, column=col, value=f"S{s['n']}")
        fec = ws.cell(row=5, column=col, value=f"{f.day}-{MESES[f.month - 1]}")
        for cc in (tit, fec):
            cc.font = Font(bold=True, color="FFFFFF", size=10)
            cc.fill = AZUL
            cc.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = 8
        tema = caps.get(s["n"])
        if tema:
            fec.comment = None
        # las clases que aun no ocurren se distinguen
        if f > hoy:
            tit.fill = PatternFill("solid", fgColor="7F9DB9")
            fec.fill = PatternFill("solid", fgColor="7F9DB9")

    col_total = 3 + len(ses)
    for r in (4, 5):
        cc = ws.cell(row=r, column=col_total, value="TOTAL" if r == 5 else "")
        cc.font = Font(bold=True, color="FFFFFF")
        cc.fill = AZUL
        cc.alignment = Alignment(horizontal="center")
    ws.column_dimensions[get_column_letter(col_total)].width = 9

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 42

    prim, ult = get_column_letter(3), get_column_letter(col_total - 1)
    for i, a in enumerate(alumnos):
        r = 6 + i
        ws.cell(row=r, column=1, value=int(a["codigo"])).number_format = "0"
        ws.cell(row=r, column=2, value=a["nombre"]).font = Font(size=10)
        for col in range(1, col_total + 1):
            ws.cell(row=r, column=col).border = BORDE
        for col in range(3, col_total):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center")
        t = ws.cell(row=r, column=col_total, value=f"=SUM({prim}{r}:{ult}{r})")
        t.font = Font(bold=True)
        t.alignment = Alignment(horizontal="center")
        if i % 2:
            for col in range(1, col_total + 1):
                if not ws.cell(row=r, column=col).fill.fgColor.rgb == "00C6EFCE":
                    ws.cell(row=r, column=col).fill = GRIS

    ultima = 5 + len(alumnos)
    tot = f"{get_column_letter(col_total)}6:{get_column_letter(col_total)}{ultima}"
    ws.conditional_formatting.add(tot, CellIsRule(operator="equal", formula=["0"], fill=ROJO))
    ws.conditional_formatting.add(tot, CellIsRule(operator="greaterThan", formula=["0"], fill=VERDE))

    # fila de resumen: cuantos participaron cada dia
    r = ultima + 2
    ws.cell(row=r, column=2, value="Participaron ese día").font = Font(bold=True, italic=True, size=9)
    for i in range(len(ses)):
        col = get_column_letter(3 + i)
        c = ws.cell(row=r, column=3 + i, value=f'=COUNTIF({col}6:{col}{ultima},">0")')
        c.font = Font(bold=True, size=9)
        c.alignment = Alignment(horizontal="center")
        c.fill = AMBAR

    ws.freeze_panes = "C6"

    salida = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "privado" / "Participacion en clase.xlsx"
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)

    print(f"[ok] {salida}")
    print(f"     {len(alumnos)} alumnos x {len(ses)} clases de {resp} antes del parcial")
    print(f"     sesiones: {', '.join(f'S{s[chr(110)]}' for s in ses)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
